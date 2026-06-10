"""
scripts/_gen_gold_brief.py
──────────────────────────
현재 DB 상태 기준 출고 검토용 GOLD 브리프 생성 (API 불필요, 순수 SQL+문서화).

분류:
  🏆 ULTRA GOLD : HIGH_CONFIRMED + fact_match EXACT/STRONG + cross_verified=CONFIRMED
  🥇 GOLD       : HIGH_CONFIRMED + fact_match EXACT/STRONG
  📰 출고 헤드라인 : n1_verified_at 존재 + newsability_score 기준 정렬

출력:
  data/exports/gold_brief_<날짜인자>.md
  data/exports/gold_brief_<날짜인자>.docx
사용: python scripts/_gen_gold_brief.py [YYYYMMDD]
"""
import sqlite3, sys, io
from pathlib import Path

try:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
except Exception:
    pass

STAMP = sys.argv[1] if len(sys.argv) > 1 else "latest"
ROOT = Path(__file__).parent.parent
DB_PATH = ROOT / "data" / "dart" / "dart_reports.db"
OUT_DIR = ROOT / "data" / "exports"
OUT_DIR.mkdir(parents=True, exist_ok=True)
MD_PATH = OUT_DIR / f"gold_brief_{STAMP}.md"
DOCX_PATH = OUT_DIR / f"gold_brief_{STAMP}.docx"

db = sqlite3.connect(str(DB_PATH), timeout=20)
db.row_factory = sqlite3.Row

GOLD_WHERE = "info_gap_label='HIGH_CONFIRMED' AND fact_match_label IN ('EXACT','STRONG')"

def rows(q):
    return db.execute(q).fetchall()

ultra = rows(f"""
  SELECT * FROM story_leads
  WHERE {GOLD_WHERE} AND cross_verified='CONFIRMED'
  ORDER BY severity DESC, corp_name
""")
gold = rows(f"""
  SELECT * FROM story_leads
  WHERE {GOLD_WHERE}
  ORDER BY severity DESC, fact_match_label, corp_name
""")
headlines = rows("""
  SELECT * FROM story_leads
  WHERE n1_verified_at IS NOT NULL AND newsability_score IS NOT NULL
  ORDER BY newsability_score DESC, publish_risk ASC
""")

def g(r, k, default=""):
    try:
        v = r[k]
    except Exception:
        return default
    return v if v not in (None, "") else default

# ── 통계 ──────────────────────────────────────────────────────────
n_total = db.execute("SELECT COUNT(*) FROM story_leads").fetchone()[0]
matrix = rows(f"""
  SELECT info_gap_label, fact_match_label, COUNT(*) c
  FROM story_leads WHERE info_gap_label LIKE 'HIGH%'
  GROUP BY 1,2 ORDER BY 1,2
""")
val_dist = rows("""
  SELECT value_direction, COUNT(*) c FROM story_leads
  WHERE value_direction IS NOT NULL GROUP BY 1 ORDER BY 2 DESC
""")

# ── Markdown ──────────────────────────────────────────────────────
def md_lead_block(r, idx):
    lines = []
    lines.append(f"#### {idx}. {g(r,'corp_name')}  ({g(r,'sector','-')})")
    hl = g(r, 'headline_ko')
    if hl:
        lines.append(f"- **출고 헤드라인:** {hl}")
    lines.append(f"- **단서:** {g(r,'title')}")
    lines.append(f"- **분류/심각도:** {g(r,'lead_type','-')} / {g(r,'severity','-')}/5")
    lines.append(f"- **검증:** fact={g(r,'fact_match_label','-')} · cross={g(r,'cross_verified','-')} · news={g(r,'news_status','-')}")
    vd = g(r, 'value_direction')
    if vd:
        lines.append(f"- **가치 방향:** {vd} · 산업대비 {g(r,'industry_alignment','-')} · 영향 {g(r,'impact_magnitude','-')}")
    ns = g(r, 'newsability_score')
    if ns != "":
        lines.append(f"- **뉴스성:** {ns}점 · valence={g(r,'valence','-')} · publish_risk={g(r,'publish_risk','-')}")
    ev = g(r, 'evidence_deep') or g(r, 'evidence')
    if ev:
        ev = ev.replace("\n", " ").strip()
        if len(ev) > 600:
            ev = ev[:600] + " …"
        lines.append(f"- **증거:** {ev}")
    rat = g(r, 'value_rationale') or g(r, 'value_rationale_ai')
    if rat:
        rat = rat.replace("\n", " ").strip()
        if len(rat) > 400:
            rat = rat[:400] + " …"
        lines.append(f"- **가치 근거:** {rat}")
    sp = g(r, 'source_path')
    if sp:
        lines.append(f"- **출처:** {sp}")
    lines.append("")
    return "\n".join(lines)

md = []
md.append(f"# 출고 검토용 GOLD 브리프 ({STAMP})\n")
md.append(f"- story_leads 총 **{n_total}건**")
md.append(f"- 🏆 ULTRA GOLD (HIGH_CONFIRMED + EXACT/STRONG + cross=CONFIRMED): **{len(ultra)}건**")
md.append(f"- 🥇 GOLD (HIGH_CONFIRMED + EXACT/STRONG): **{len(gold)}건**")
md.append(f"- 📰 출고 헤드라인 생성분(_n1): **{len(headlines)}건**\n")

md.append("## 검증 매트릭스 (info_gap × fact_match)\n")
md.append("| info_gap | fact_match | 건수 |")
md.append("|---|---|---|")
for r in matrix:
    md.append(f"| {r['info_gap_label']} | {r['fact_match_label']} | {r['c']} |")
md.append("")

if val_dist:
    md.append("## 가치 방향 분포\n")
    md.append("| 방향 | 건수 |")
    md.append("|---|---|")
    for r in val_dist:
        md.append(f"| {r['value_direction']} | {r['c']} |")
    md.append("")

md.append("---\n")
md.append(f"## 🏆 ULTRA GOLD ({len(ultra)}건) — 최우선 출고 후보\n")
if ultra:
    for i, r in enumerate(ultra, 1):
        md.append(md_lead_block(r, i))
else:
    md.append("_해당 없음_\n")

md.append("---\n")
md.append(f"## 📰 출고 헤드라인 (newsability 순, {len(headlines)}건)\n")
if headlines:
    md.append("| 점수 | valence/risk | 회사 | 헤드라인 |")
    md.append("|---|---|---|---|")
    for r in headlines:
        md.append(f"| {g(r,'newsability_score','-')} | {g(r,'valence','-')}/{g(r,'publish_risk','-')} | {g(r,'corp_name')} | {g(r,'headline_ko','-')} |")
    md.append("")
else:
    md.append("_아직 _n1 미생성 (quota 회복 후 생성 예정)_\n")

md.append("---\n")
md.append(f"## 🥇 GOLD 일람 ({len(gold)}건)\n")
md.append("| # | 회사 | 섹터 | 심각도 | fact | cross | news | 단서 |")
md.append("|---|---|---|---|---|---|---|---|")
for i, r in enumerate(gold, 1):
    title = g(r, 'title').replace("|", "/")
    if len(title) > 50:
        title = title[:50] + "…"
    md.append(f"| {i} | {g(r,'corp_name')} | {g(r,'sector','-')} | {g(r,'severity','-')} | "
              f"{g(r,'fact_match_label','-')} | {g(r,'cross_verified','-')} | {g(r,'news_status','-')} | {title} |")
md.append("")

MD_PATH.write_text("\n".join(md), encoding="utf-8")
print(f"[MD] {MD_PATH}  ({len(md)} 줄)")

# ── Word ──────────────────────────────────────────────────────────
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    st = doc.styles['Normal']; st.font.name = '맑은 고딕'; st.font.size = Pt(10)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    def para(text, size=10, bold=False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = '맑은 고딕'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        run.font.size = Pt(size); run.bold = bold
        if color: run.font.color.rgb = color
        return p

    para(f"출고 검토용 GOLD 브리프 ({STAMP})", 16, True)
    para(f"story_leads 총 {n_total}건  |  ULTRA GOLD {len(ultra)}  |  GOLD {len(gold)}  |  헤드라인 {len(headlines)}", 10)

    para(f"🏆 ULTRA GOLD ({len(ultra)}건)", 14, True, RGBColor(0xC0, 0x00, 0x00))
    for i, r in enumerate(ultra, 1):
        para(f"{i}. {g(r,'corp_name')} ({g(r,'sector','-')})", 11, True)
        if g(r, 'headline_ko'):
            para(f"   헤드라인: {g(r,'headline_ko')}", 10, True)
        para(f"   단서: {g(r,'title')}", 10)
        para(f"   검증: fact={g(r,'fact_match_label','-')} · cross={g(r,'cross_verified','-')} · news={g(r,'news_status','-')}", 9)
        ev = (g(r, 'evidence_deep') or g(r, 'evidence') or "").replace("\n", " ")
        if ev:
            para(f"   증거: {ev[:800]}", 9)

    para(f"📰 출고 헤드라인 ({len(headlines)}건)", 14, True, RGBColor(0x00, 0x55, 0xAA))
    for r in headlines:
        para(f"[{g(r,'newsability_score','-')}점/{g(r,'publish_risk','-')}] {g(r,'corp_name')} — {g(r,'headline_ko','-')}", 10)

    para(f"🥇 GOLD 일람 ({len(gold)}건)", 14, True)
    tbl = doc.add_table(rows=1, cols=6); tbl.style = 'Light Grid Accent 1'
    hdr = tbl.rows[0].cells
    for c, t in zip(hdr, ["회사", "섹터", "심각도", "fact", "cross", "단서"]):
        c.text = t
    for r in gold:
        cells = tbl.add_row().cells
        cells[0].text = str(g(r, 'corp_name'))
        cells[1].text = str(g(r, 'sector', '-'))
        cells[2].text = str(g(r, 'severity', '-'))
        cells[3].text = str(g(r, 'fact_match_label', '-'))
        cells[4].text = str(g(r, 'cross_verified', '-'))
        cells[5].text = str(g(r, 'title'))[:60]

    doc.save(str(DOCX_PATH))
    print(f"[DOCX] {DOCX_PATH}")
except Exception as e:
    print(f"[DOCX] 생성 실패 (MD는 정상): {e}")

db.close()
print("완료.")
